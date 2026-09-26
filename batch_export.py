# -*- coding: utf-8 -*-
"""批量问答报告导出 —— batch_export.py（47 号方案阶段 B）

两个出口、两种读者：
  - **DOCX**：给面试官看的报告（可读性优先，含边界说明）；
  - **CSV**：给机器看的明细（UTF-8 BOM，Excel 直接打开不乱码）。

三条必须守住的规矩（47 号方案 §4）：
  1. **只读本次实际记录**：入参是 batch_jobs 的任务快照，这里不检索、不生成、不发任何请求；
  2. **全部在内存里生成**：返回完整字节，由 api_v3 一次性响应。
     这比"写临时文件再返回"更不容易留下半截文件（生成失败时根本没有文件产生）；
  3. **不写凭据、不写绝对路径、不写内部栈**：错误列只放 batch_jobs 已经 scrub 过的短文本；
     用户文本（问题、来源标题）一律按普通文本写，不参与任何路径拼接，也不做 HTML/Markdown 解释。

CSV 的状态列用**稳定英文状态码**（与任务接口一致）：done / refused / failed / not_executed。
DOCX 里用中文标签，并在汇总处同时给出计划/实际模型请求数。

50 号工单新增（2026-09-22，报告缺中文字体声明 + 命中证据未落地）：
  - CSV 表尾追加 `hit_chunk_ids` / `hit_source_titles` / `hit_sections` / `hit_scores` /
    `hit_snippets` 五个**命中证据列**，取值来自 `item["hits"]`（批量检索时已连同原文一起
    保存），顺序与该数组一致；`citation_ids` / `source_titles` 仍然只表示"回答里解析出的
    引用"，仅检索模式下为空数组 —— 两类"引用"分列存放，互不冒充；
  - DOCX 逐题把命中改成"一条命中一个区块"（排名 / chunk id / 来源标题 / 章节 / BM25 分数 /
    原文摘录，摘录上限 240 字并标注截断，完整文本在 CSV 的 `hit_snippets`）；
    仅检索模式措辞改为「回答：未生成（仅检索模式）」「回答引用：不适用（未生成回答）」；
  - 导出统一写死中文字体：docDefaults + Normal/Title/Heading 1/Heading 2 + 每个 run 兜底，
    并清掉主题字体引用，`themeFontLang@eastAsia` 设为 `zh-CN`（不再依赖主题里为空的中文链）。
"""
from __future__ import annotations

import csv
import io
import json
import time

APP_VERSION = "0.5.2"          # 与 launcher.py / build_release.py 的产品版本同步

# CSV 前 10 列语义不变（citation_ids/source_titles 继续表示**回答中实际解析出的引用**，
# 仅检索模式下就是空数组）；第 11–15 列是 50 号工单新增的**独立命中证据列**，
# 一律取 item["hits"] 已保存的快照，顺序与该数组完全一致。
CSV_HEADERS = ("sequence", "question", "status", "answer", "citation_ids",
               "source_titles", "error", "mode", "knowledge_base", "top_k",
               "hit_chunk_ids", "hit_source_titles", "hit_sections",
               "hit_scores", "hit_snippets")

# DOCX 里的命中原文摘录上限（CSV 的 hit_snippets 不截断，保留完整已保存文本）
SNIPPET_MAX_CHARS = 240

# 导出统一使用的中文字体（50 号工单第四节）：显式写名字，不留给主题默认链。
# 主题链里的 `<a:ea typeface=""/>` 是空值、`themeFontLang@eastAsia` 又是 ja-JP，
# 于是"用哪个中文字体"只能由渲染器各自猜 —— 部分汉字就渲染成方框（豆腐块）。
CJK_FONT = "Microsoft YaHei"   # 微软雅黑
LATIN_FONT = "Calibri"
EAST_ASIA_LANG = "zh-CN"

DOCX_BOUNDARY = (
    "边界说明（固定文案）：引用有效不等于回答必然正确；本报告的生成质量结论受样本量与模型"
    "随机性限制，单轮结果不构成对生成层的正式验收。检索命中率（Hit@K / MRR）是检索口径，"
    "不是答案正确率，且只适用于内置学习库——「我的资料库」不继承该指标。"
    "本报告由本机工具生成，不是公网服务，不用于电力现场安全决策。"
)

MODE_LABEL = {"retrieve_only": "仅检索（不调用模型）", "generate": "检索并生成"}
STATUS_LABEL = {
    "pending": "等待",
    "running": "运行中",
    "done": "完成",
    "refused": "未找到资料依据（模型按 prompt 明确拒答）",
    "failed": "失败",
    "not_executed": "未执行（任务被取消或提前停止）",
}


def export_filename(job: dict, fmt: str) -> str:
    """安全文件名：只用固定前缀 + 时间戳，**不含任何用户文本**。"""
    stamp = time.strftime("%Y%m%d-%H%M%S")
    return f"RAG批量问答报告-{stamp}.{'csv' if fmt == 'csv' else 'docx'}"


def _summary_line(job: dict) -> dict:
    from batch_jobs import summary                      # 单一口径，避免两处各算一套
    return summary(job)


# ---------------- CSV ----------------
def render_csv(job: dict) -> bytes:
    """UTF-8 BOM + 标准转义；多值字段用稳定 JSON 字符串，顺序与任务快照一致。

    两类 "引用" 必须分列存放，不许互相冒充：
      - `citation_ids` / `source_titles`：**回答里实际解析出的引用**（仅检索模式为空数组）；
      - `hit_*`：这次检索**真实命中**的前 N 条（含原文），在任何模式下都必须落地。
    """
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(CSV_HEADERS)
    for item in job["items"]:
        refs = item.get("refs") or []
        hits = item.get("hits") or []
        writer.writerow([
            item["sequence"],
            item["question"],
            item["status"],
            item.get("answer") or "",
            json.dumps([r["id"] for r in refs], ensure_ascii=False),
            json.dumps([r["source_title"] for r in refs], ensure_ascii=False),
            item.get("error") or "",
            job["mode"],
            job["corpus_name"],
            job["top_k"],
            json.dumps([h.get("id") or "" for h in hits], ensure_ascii=False),
            json.dumps([h.get("source_title") or "" for h in hits], ensure_ascii=False),
            json.dumps([h.get("section") or "" for h in hits], ensure_ascii=False),
            json.dumps([h.get("score") for h in hits], ensure_ascii=False),
            json.dumps([h.get("text") or "" for h in hits], ensure_ascii=False),
        ])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


# ---------------- DOCX 辅助：字体与逐题区块 ----------------
def _set_rfonts(rpr, latin: str = LATIN_FONT, east_asian: str = CJK_FONT) -> None:
    """给一个 `w:rPr` 写**显式字体名**，并清掉主题引用。

    为什么必须连主题引用一起清：只要 `w:asciiTheme` / `w:eastAsiaTheme` 还在，这段文字
    就仍然"由主题决定字体"；而 python-docx 默认主题里 `<a:ea typeface=""/>` 是空值 ——
    等于整篇都没有"该用哪个中文字体"的答案，只能由渲染器各自猜（猜成日文字体就出方框）。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        # 插到 w:rStyle 之后（OOXML 对 rPr 子元素有顺序要求）
        rstyle = rpr.find(qn("w:rStyle"))
        if rstyle is not None:
            rstyle.addnext(rfonts)
        else:
            rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:cs"), latin)
    rfonts.set(qn("w:eastAsia"), east_asian)


def _apply_cjk_fonts(doc) -> None:
    """三层落地显式中文字体（50 号工单第四节）：默认链 → 命名样式 → 每个 run 兜底。

    不嵌字体、不重写 theme1.xml：只保证"已经用到的样式和正文 run 都写明了东亚字体"。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    styles_el = doc.styles.element

    # 1) docDefaults：整篇的默认链（原来只有 theme 引用）
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles_el.insert(0, doc_defaults)
    rpr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        doc_defaults.insert(0, rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    _set_rfonts(rpr)

    # 2) 实际用到的命名样式（标题都出自这几个）
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        _set_rfonts(style.element.get_or_add_rPr())

    # 3) 正文 runs 兜底：任何一处漏掉样式继承都不会再退回空主题字体
    for para in doc.paragraphs:
        for run in para.runs:
            _set_rfonts(run._element.get_or_add_rPr())

    # 4) 东亚语言：ja-JP → zh-CN（这条正是"方框只出现在部分汉字上"的直接原因）
    settings_el = doc.settings.element
    tfl = settings_el.find(qn("w:themeFontLang"))
    if tfl is None:
        tfl = OxmlElement("w:themeFontLang")
        settings_el.append(tfl)
    tfl.set(qn("w:val"), "en-US")
    tfl.set(qn("w:eastAsia"), EAST_ASIA_LANG)


def _add_answer_block(doc, job: dict, item: dict) -> None:
    """回答 / 回答引用：把"未生成"和"未命中"分开说，不再用含糊的"无回答文本 / 引用：无"。"""
    answer = item.get("answer")
    notice = item.get("notice")
    if answer:
        doc.add_paragraph("回答：" + answer)
    elif job["mode"] == "retrieve_only":
        doc.add_paragraph("回答：未生成（仅检索模式）")
    elif notice == "zero_hits":
        doc.add_paragraph("回答：未生成（检索结果为 0 条，已短路，未调用模型）")
    elif notice == "generation_failed":
        doc.add_paragraph("回答：未生成（生成失败，原因见下方错误摘要）")
    else:
        doc.add_paragraph("回答：未生成（该题未执行）")

    refs = item.get("refs") or []
    if refs:
        doc.add_paragraph(f"回答引用（{len(refs)} 个，去重后）：")
        for r in refs:
            doc.add_paragraph(f"[{r['id']}] {r['source_title']} · {r['section']}"
                              + ("（同节补充上下文）" if r.get("origin") == "neighbor" else ""))
            doc.add_paragraph(r["text"])
    elif answer:
        doc.add_paragraph("回答引用：无（模型回答里没有出现任何引用标记）")
    else:
        doc.add_paragraph("回答引用：不适用（未生成回答）")


def _add_hits_block(doc, item: dict) -> None:
    """检索命中：**一条命中一个独立区块**，展示排名 / chunk id / 来源 / 章节 / 分数 / 原文。"""
    hits = item.get("hits") or []
    if not hits:
        if item["status"] in ("pending", "running", "not_executed"):
            doc.add_paragraph("检索命中：不适用（该题未执行，没有检索记录）")
        else:
            doc.add_paragraph("检索命中：0 条（本次检索没有返回任何片段）")
        return
    doc.add_paragraph(f"检索命中（前 {len(hits)} 条，按 BM25 排名）：")
    for rank, h in enumerate(hits, 1):
        doc.add_paragraph(f"命中 {rank}／{len(hits)}：{h.get('id') or '（无 id）'}")
        doc.add_paragraph(f"来源标题：{h.get('source_title') or '（无）'}")
        doc.add_paragraph(f"章节：{h.get('section') or '（无章节标题）'}")
        score = h.get("score")
        doc.add_paragraph(f"BM25 分数：{score if score is not None else '（无）'}")
        text = h.get("text") or ""
        if not text:
            doc.add_paragraph("命中原文：（该任务快照未保存原文，导出侧不补检索）")
        elif len(text) > SNIPPET_MAX_CHARS:
            doc.add_paragraph(f"命中原文（原文共 {len(text)} 字，此处截断显示前 "
                              f"{SNIPPET_MAX_CHARS} 字；完整文本见 CSV 的 hit_snippets 列）：")
            doc.add_paragraph(text[:SNIPPET_MAX_CHARS] + "…")
        else:
            doc.add_paragraph("命中原文：")
            doc.add_paragraph(text)


# ---------------- DOCX ----------------
def render_docx(job: dict) -> bytes:
    """面试官可读报告（47 号方案 B1 的 11 项，顺序即章节顺序）。"""
    from docx import Document

    s = _summary_line(job)
    snap = job.get("corpus_snapshot") or {}
    doc = Document()

    doc.add_heading("RAG 批量问答报告", 0)
    doc.add_paragraph(f"生成时间：{time.strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"项目版本：{APP_VERSION}")

    doc.add_heading("一、知识库", level=1)
    kb = [f"名称：{job['corpus_name']}（{job['corpus_id']}）"]
    if snap:
        kb.append(f"来源数：{snap.get('source_count', '未知')}；"
                  f"片段数：{snap.get('chunk_count', '未知')}；"
                  f"总字符数：{snap.get('total_chars', '未知')}")
        kb.append(f"知识库版本：{snap.get('version', '未知')}；"
                  f"指纹：{(snap.get('fingerprint') or '未知')[:16]}")
    else:
        kb.append("本次任务未记录知识库统计快照。")
    for line in kb:
        doc.add_paragraph(line)

    doc.add_heading("二、本次运行", level=1)
    doc.add_paragraph(f"运行模式：{MODE_LABEL.get(job['mode'], job['mode'])}")
    doc.add_paragraph(f"Top-K：{job['top_k']}")
    doc.add_paragraph(f"问题总数：{job['total']}")
    doc.add_paragraph(f"计划模型请求数：{job['planned_calls']}；"
                      f"实际启动：{job['started_calls']}；成功：{s['succeeded']}；"
                      f"自动重试：{job['retries']}")
    doc.add_paragraph("生成目标：" + ("本机回环地址（未调用真实模型供应商）"
                                    if job.get("llm_loopback") else "外部模型服务")
                      + f"；模型：{job.get('llm_model') or '未使用'}")
    doc.add_paragraph(f"任务状态：{job['state']}"
                      + (f"；原因：{job['stop_reason']}" if job.get("stop_reason") else ""))

    doc.add_heading("三、汇总", level=1)
    doc.add_paragraph(f"完成：{s['done']}；拒答（未找到资料依据）：{s['refused']}；"
                      f"失败：{s['failed']}；未执行：{s['not_executed']}")
    doc.add_paragraph(DOCX_BOUNDARY)

    doc.add_heading("四、逐题明细", level=1)
    for item in job["items"]:
        head = doc.add_heading(f"第 {item['sequence']} 题", level=2)
        head.paragraph_format.keep_with_next = True      # 页尾不孤悬一个"第 N 题"
        doc.add_paragraph(f"状态：{STATUS_LABEL.get(item['status'], item['status'])}")
        qp = doc.add_paragraph(f"问题：{item['question']}")
        qp.paragraph_format.keep_with_next = True        # 问题不孤悬页尾
        _add_answer_block(doc, job, item)
        _add_hits_block(doc, item)
        if item.get("error"):
            doc.add_paragraph(f"错误摘要：{item['error']}")

    _apply_cjk_fonts(doc)        # 内容写完后统一兜底：显式东亚字体 + zh-CN
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def build_export(job: dict, fmt: str) -> dict:
    """统一出口：{filename, content_type, data}。fmt 只允许 csv / docx。"""
    if fmt == "csv":
        return {"filename": export_filename(job, "csv"),
                "content_type": "text/csv; charset=utf-8",
                "data": render_csv(job)}
    if fmt == "docx":
        return {"filename": export_filename(job, "docx"),
                "content_type": ("application/vnd.openxmlformats-officedocument"
                                 ".wordprocessingml.document"),
                "data": render_docx(job)}
    raise ValueError(f"不支持的导出格式：{fmt!r}（只允许 csv / docx）")
