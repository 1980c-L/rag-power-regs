# -*- coding: utf-8 -*-
"""批量提问 + 报告导出验收（47 号方案阶段 A/B 的服务端部分）—— verify_batch_jobs.py

它验什么（判据全部来自 47 号方案 §3/§4/§7，不新增也不放宽）：
  1. 输入硬上限与 fail closed：1–20 条、单条 ≤ 500 字符、空行过滤、重复保留、未知库/模式/Top-K；
  2. 仅检索模式的模型请求数**严格为 0**；生成模式的计划数 = 实际启动数、零重试；
  3. **并发固定为 1**：用专用桩的 in_flight 峰值证明（不是靠"看起来是顺序的"）；
  4. 系统性错误（500 / 401）后不再启动后续；非系统性错误（返回体结构异常）只记这一条失败；
  5. 取消语义：当前这条跑完并计入实际调用数，之后不再启动；**并且**（49 号 P1）
     在"当前题仍在检索、模型请求尚未发出"的窗口里取消，也必须 0 请求；
  6. 同一进程只允许一个生成任务；重启后旧任务不恢复、不重放；
  7. DOCX/CSV 导出：逐字段对拍、零新增调用、注入/穿越反例、无凭据与绝对路径；
  8. 薄入口的 HTTP 契约（api_v3）：字段校验、状态码映射、下载响应头。

**全程零真实模型请求**：生成一律指向 `tools/batch_stub_server.py`（本机回环 + 行为脚本）。
本文件内不写任何个人绝对路径；批量状态目录与用户库目录都被隔离到系统临时目录。

运行：
    python tests/verify_batch_jobs.py [--out <证据目录>] [--keep]
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

# ---------------- 测试隔离：必须在 import batch_jobs 之前设置环境变量 ----------------
WORK = Path(tempfile.mkdtemp(prefix="rag_batch_verify_"))
os.environ["RAG_BATCH_DIR"] = str(WORK / "batches")
os.environ["RAG_USER_LIBRARY_DIR"] = str(WORK / "user_library")   # 空库（409 用例）

import rag_core as rc                                             # noqa: E402
import batch_jobs as bj                                           # noqa: E402
import batch_export as bx                                         # noqa: E402

STUB_PORT = 5291
API_PORT = 5290
STUB_STATE = f"http://127.0.0.1:{STUB_PORT}/__batch_stub__/state"
STUB_RESET = f"http://127.0.0.1:{STUB_PORT}/__batch_stub__/reset"
STUB_BASE = f"http://127.0.0.1:{STUB_PORT}/v1"
STUB_LOG = WORK / "batch_stub_requests.jsonl"
API = f"http://127.0.0.1:{API_PORT}"
SENTINEL_KEY = "sk-BATCH-SENTINEL-DO-NOT-LEAK-7c21"

Q_TRIAD = "RAG 评估三元组包含哪三个维度？"
Q_GROUND = "装设接地线的顺序是什么？"
Q_FLOW = "RAG 的基本流程是什么？"
Q_ZERO = "红烧肉怎么做？"

PASSED: list = []
FAILED: list = []
PROCESSES: list = []
EVIDENCE: dict = {"phases": []}
_EXPORT_JOB_IDS: dict = {}
_EXTRA_FILES: dict = {}       # {文件名: bytes}：样本文件由 main 统一落到证据目录
MANIFEST = REPO / "data" / "rag_learning" / "corpus_manifest.json"
MANIFEST_HASH_BEFORE = ""


def check(name: str, ok: bool, detail: str = "") -> None:
    (PASSED if ok else FAILED).append(name)
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""), flush=True)


# ---------------- 桩 ----------------
def stub_state() -> dict:
    with urllib.request.urlopen(STUB_STATE, timeout=10) as r:
        return json.loads(r.read().decode("utf-8"))


def stub_reset(script: str = "", delay: float | None = None) -> None:
    params: dict = {}
    if script:
        params["script"] = script
    if delay is not None:
        params["delay"] = str(delay)
    url = STUB_RESET + ("?" + urllib.parse.urlencode(params) if params else "")
    with urllib.request.urlopen(url, timeout=10) as r:
        r.read()


def stub_log() -> list:
    if not STUB_LOG.exists():
        return []
    return [json.loads(line) for line in STUB_LOG.read_text(encoding="utf-8").splitlines() if line]


def start_stub(script: str = "", delay: float = 0.02) -> subprocess.Popen:
    proc = subprocess.Popen(
        [sys.executable, "tools/batch_stub_server.py", "--port", str(STUB_PORT),
         "--log", str(STUB_LOG), "--script", script, "--delay", str(delay)],
        cwd=str(REPO), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    PROCESSES.append(proc)
    for _ in range(80):
        try:
            stub_state()
            return proc
        except Exception:                                          # noqa: BLE001
            time.sleep(0.25)
    raise RuntimeError("批量桩没有起来")


# ---------------- api_v3 子进程 ----------------
def start_api(with_key: bool) -> subprocess.Popen:
    env = dict(os.environ)
    env["PYTHONIOENCODING"] = "utf-8"
    if with_key:
        env["DEEPSEEK_API_KEY"] = SENTINEL_KEY
    else:
        env.pop("DEEPSEEK_API_KEY", None)
    log = WORK / ("api_with_key.log" if with_key else "api_no_key.log")
    handle = log.open("w", encoding="utf-8")
    PROCESSES.append(handle)
    proc = subprocess.Popen([sys.executable, "api_v3.py", "--port", str(API_PORT),
                             "--llm-base-url", STUB_BASE, "--llm-model", "batch-stub"],
                            cwd=str(REPO), env=env, stdout=handle, stderr=subprocess.STDOUT)
    PROCESSES.append(proc)
    for _ in range(80):
        try:
            code, _, _, _ = http("GET", "/api/health")
            if code == 200:
                return proc
        except Exception:                                          # noqa: BLE001
            time.sleep(0.25)
    raise RuntimeError("api_v3 没有起来")


def stop_process(proc) -> None:
    if proc is None:
        return
    if isinstance(proc, subprocess.Popen):
        if os.name == "nt":
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        else:
            proc.terminate()
        try:
            proc.wait(timeout=15)
        except Exception:                                          # noqa: BLE001
            proc.kill()


def http(method: str, path: str, payload=None, timeout: int = 60):
    """返回 (status, headers, body_text, body_bytes)。"""
    url = f"{API}{path}"
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8") if payload is not None else None
    headers = {"Content-Type": "application/json"} if data is not None else {}
    req = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            raw = r.read()
            return r.status, dict(r.headers), raw.decode("utf-8", "replace"), raw
    except urllib.error.HTTPError as e:
        raw = e.read()
        return e.code, dict(e.headers), raw.decode("utf-8", "replace"), raw


def jbody(text: str) -> dict:
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return {}


# ---------------- 任务工具 ----------------
def create(**kw) -> dict:
    kw.setdefault("corpus_id", rc.RAG_LEARNING)
    kw.setdefault("mode", "retrieve_only")
    kw.setdefault("top_k", 3)
    kw.setdefault("llm_base_url", STUB_BASE)
    kw.setdefault("llm_model", "batch-stub")
    kw.setdefault("api_key", SENTINEL_KEY if kw["mode"] == "generate" else "")
    if "confirm_calls" not in kw:
        kw["confirm_calls"] = len(kw.get("questions") or []) if kw["mode"] == "generate" else 0
    return bj.create_job(**kw)


def wait_job(job_id: str, timeout: float = 180.0) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = bj.get_job(job_id)
        if job["state"] in bj.FINAL_STATES:
            return job
        time.sleep(0.05)
    raise AssertionError(f"任务 {job_id} 在 {timeout}s 内没有结束"
                         f"（state={bj.get_job(job_id)['state']}）")


def wait_until(pred, timeout: float = 10.0) -> bool:
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            if pred():
                return True
        except Exception:                                          # noqa: BLE001
            pass
        time.sleep(0.05)
    return False


# ---------------- 阶段 1：输入校验 ----------------
def phase_validation() -> None:
    stub_reset()
    base = dict(corpus_id=rc.RAG_LEARNING, mode="retrieve_only", top_k=3, confirm_calls=0,
                llm_base_url=STUB_BASE, llm_model="batch-stub", api_key="")

    def expect_code(name: str, code: str, fn) -> None:
        try:
            fn()
            check(name, False, "预期被拒绝，实际通过了")
        except bj.BatchError as exc:
            check(name, exc.code == code, f"实际 {exc.code}（{exc.message}）")
        except Exception as exc:                                    # noqa: BLE001
            check(name, False, f"抛了非 BatchError：{type(exc).__name__}: {exc}")

    expect_code("批量·校验：21 条被拒（too_many_questions）", "too_many_questions",
                lambda: bj.create_job(questions=[f"q{i}" for i in range(21)], **base))
    expect_code("批量·校验：全是空行被拒（empty_questions）", "empty_questions",
                lambda: bj.create_job(questions=["", "   ", "\n"], **base))
    expect_code("批量·校验：单条超过 500 字符被拒（question_too_long）", "question_too_long",
                lambda: bj.create_job(questions=["x" * 501], **base))
    expect_code("批量·校验：非字符串元素被拒（invalid_question）", "invalid_question",
                lambda: bj.create_job(questions=["正常问题", 3], **base))
    expect_code("批量·校验：questions 不是数组被拒（invalid_questions）", "invalid_questions",
                lambda: bj.create_job(questions="a\nb", **base))
    expect_code("批量·校验：未知知识库被拒（unknown_corpus）", "unknown_corpus",
                lambda: bj.create_job(questions=["a"], **{**base, "corpus_id": "nope"}))
    expect_code("批量·校验：未知模式被拒（unknown_mode）", "unknown_mode",
                lambda: bj.create_job(questions=["a"], **{**base, "mode": "auto"}))
    expect_code("批量·校验：Top-K 越界被拒（top_k_out_of_range）", "top_k_out_of_range",
                lambda: bj.create_job(questions=["a"], **{**base, "top_k": 6}))
    expect_code("批量·校验：Top-K 为布尔被拒（invalid_top_k）", "invalid_top_k",
                lambda: bj.create_job(questions=["a"], **{**base, "top_k": True}))
    expect_code("批量·校验：生成模式无 Key 整体拒绝（no_api_key）", "no_api_key",
                lambda: bj.create_job(mode="generate", questions=["a"], confirm_calls=1,
                                      api_key="", corpus_id=rc.RAG_LEARNING, top_k=3,
                                      llm_base_url=STUB_BASE, llm_model="batch-stub"))
    expect_code("批量·校验：生成模式确认值与计划值不一致被拒（confirm_mismatch）", "confirm_mismatch",
                lambda: bj.create_job(mode="generate", questions=["a", "b"], confirm_calls=1,
                                      api_key=SENTINEL_KEY, corpus_id=rc.RAG_LEARNING, top_k=3,
                                      llm_base_url=STUB_BASE, llm_model="batch-stub"))
    expect_code("批量·校验：仅检索模式带确认值被拒（confirm_mismatch）", "confirm_mismatch",
                lambda: bj.create_job(questions=["a"], **{**base, "confirm_calls": 3}))
    check("批量·校验：以上全部拒绝都没有发出任何生成请求（桩计数 0）",
          stub_state()["count"] == 0, f"桩计数={stub_state()['count']}")
    EVIDENCE["phases"].append({"name": "validation", "stub_requests": stub_state()["count"]})


# ---------------- 阶段 2：仅检索 ----------------
def phase_retrieve() -> None:
    stub_reset()
    job = wait_job(create(questions=[Q_TRIAD], mode="retrieve_only")["job_id"])
    exp = rc.expand_and_assemble(Q_TRIAD, top_k=3, corpus_id=rc.RAG_LEARNING)
    item = job["items"][0]
    check("批量·仅检索：任务 COMPLETED、单条 done、notice=retrieve_only",
          job["state"] == "COMPLETED" and item["status"] == "done"
          and item["notice"] == "retrieve_only",
          f"state={job['state']} item={item['status']}/{item['notice']}")
    check("批量·仅检索：命中 id 与顺序与直接调 rag_core 完全一致",
          [h["id"] for h in item["hits"]] == [e["id"] for e in exp["hits"]],
          f"api={[h['id'] for h in item['hits']]} rc={[e['id'] for e in exp['hits']]}")
    check("批量·仅检索：BM25 分数逐字段一致（不做二次四舍五入）",
          [h["score"] for h in item["hits"]] == [e["score"] for e in exp["hits"]])
    check("批量·仅检索：没有回答、没有引用（检索模式不生成）",
          item["answer"] is None and item["refs"] == [])
    check("批量·仅检索：计划数/实际启动数都是 0，桩计数为 0",
          job["planned_calls"] == 0 and job["started_calls"] == 0
          and stub_state()["count"] == 0, f"桩计数={stub_state()['count']}")

    qs20 = [f"第 {i} 个问题：{Q_TRIAD}" for i in range(1, 21)]
    job20 = wait_job(create(questions=qs20)["job_id"])
    check("批量·仅检索：20 条全部完成且严格保持输入顺序",
          job20["total"] == 20 and [i["question"] for i in job20["items"]] == qs20
          and all(i["status"] == "done" for i in job20["items"]),
          f"total={job20['total']}")
    check("批量·仅检索：20 条也没有发出任何生成请求", stub_state()["count"] == 0)

    job_mix = wait_job(
        create(questions=["同一个问题？", "同一个问题？", "", "   ", "另一个问题？"])["job_id"])
    check("批量·过滤：空行/纯空白被丢弃（total=3）", job_mix["total"] == 3,
          f"total={job_mix['total']}")
    check("批量·重复：相同问题不静默去重，两条都单独执行",
          job_mix["items"][0]["question"] == job_mix["items"][1]["question"]
          and job_mix["items"][0]["status"] == "done" and job_mix["items"][1]["status"] == "done"
          and job_mix["items"][0]["sequence"] != job_mix["items"][1]["sequence"])
    EVIDENCE["phases"].append({"name": "retrieve", "stub_requests": stub_state()["count"]})


# ---------------- 阶段 3：生成 ----------------
def phase_generate() -> None:
    stub_reset(delay=0.02)
    qs = [Q_TRIAD, Q_GROUND, Q_FLOW]
    job = wait_job(create(questions=qs, mode="generate")["job_id"])
    st = stub_state()
    check("批量·生成：计划数 = 实际启动数 = 成功数（每题最多 1 次）",
          job["planned_calls"] == 3 and job["started_calls"] == 3 and job["succeeded"] == 3,
          f"planned={job['planned_calls']} started={job['started_calls']} ok={job['succeeded']}")
    check("批量·生成：桩恰好收到 3 次请求（零重试、无额外调用）", st["count"] == 3,
          f"桩计数={st['count']}")
    check("批量·生成：并发峰值为 1（串行是结构性事实，不是看起来像）",
          st["max_in_flight"] == 1, f"max_in_flight={st['max_in_flight']}")
    check("批量·生成：retries 恒为 0（没有重试循环）", job["retries"] == 0)

    log = stub_log()
    ok_refs = True
    for i, item in enumerate(job["items"]):
        record = log[i] if i < len(log) else {}
        want = rc.extract_refs(record.get("answer", ""), set(record.get("ids_in_prompt", [])))
        if [r["id"] for r in item["refs"]] != want:
            ok_refs = False
    check("批量·生成：引用与桩日志里的回答原文独立复算一致（不靠自证）", ok_refs,
          f"每题引用数={[len(i['refs']) for i in job['items']]}")
    check("批量·生成：引用条目带来源标题与证据片段（导出要用）",
          all(all(r.get("source_title") and r.get("text") for r in i["refs"])
              for i in job["items"] if i["refs"]))

    stub_reset()
    check("批量·生成：零命中用例的前提成立（该问题在电力示例库中确实零命中）",
          rc.retrieve(Q_ZERO, top_k=3, corpus_id=rc.POWER_DEMO) == [],
          "前置条件不成立时，下面的零命中判据本身就没有意义")
    job_zero = wait_job(create(questions=[Q_ZERO], mode="generate",
                               corpus_id=rc.POWER_DEMO)["job_id"])
    check("批量·生成：零命中题不调用模型（桩计数 0）", stub_state()["count"] == 0,
          f"桩计数={stub_state()['count']}")
    check("批量·生成：零命中题 notice=zero_hits、无回答、状态照常 done",
          job_zero["items"][0]["notice"] == "zero_hits"
          and job_zero["items"][0]["answer"] is None
          and job_zero["items"][0]["status"] == "done",
          f"notice={job_zero['items'][0]['notice']} hits={len(job_zero['items'][0]['hits'])}")

    stub_reset(delay=0.01)
    qs20 = ([Q_TRIAD, Q_GROUND, Q_FLOW] * 7)[:20]
    job20 = wait_job(create(questions=qs20, mode="generate")["job_id"], timeout=300)
    st20 = stub_state()
    check("批量·生成：20 题实际启动数恰好 20，硬上限未被突破",
          job20["started_calls"] == 20 and job20["planned_calls"] == 20,
          f"started={job20['started_calls']}")
    check("批量·生成：20 题桩计数恰好 20（0 重试、0 额外请求）", st20["count"] == 20,
          f"桩计数={st20['count']}")
    check("批量·生成：20 题并发峰值仍为 1", st20["max_in_flight"] == 1,
          f"max_in_flight={st20['max_in_flight']}")
    EVIDENCE["phases"].append({"name": "generate", "stub_requests": st20["count"],
                               "max_in_flight": st20["max_in_flight"],
                               "job20_started": job20["started_calls"]})


# ---------------- 阶段 4：系统性错误 / 非系统性错误 / 取消 ----------------
def phase_errors_and_cancel() -> None:
    stub_reset(script="ok,http500,ok,ok")
    job = wait_job(create(questions=[Q_TRIAD, Q_GROUND, Q_FLOW, Q_TRIAD],
                          mode="generate")["job_id"])
    check("批量·停手：第 2 条遇到 500 后任务为 STOPPED_ON_ERROR 且写明原因",
          job["state"] == "STOPPED_ON_ERROR" and bool(job["stop_reason"]),
          f"state={job['state']} reason={job['stop_reason']}")
    check("批量·停手：实际启动数为 2（不再启动后续问题）", job["started_calls"] == 2,
          f"started={job['started_calls']}")
    check("批量·停手：桩只收到 2 次请求（同一个系统错误没有被重复消费）",
          stub_state()["count"] == 2, f"桩计数={stub_state()['count']}")
    check("批量·停手：出错的那一条如实记为 failed 且带错误摘要",
          job["items"][1]["status"] == "failed" and bool(job["items"][1]["error"]))
    check("批量·停手：未执行的两条记为 not_executed（不伪装成失败或空答案）",
          [i["status"] for i in job["items"][2:]] == ["not_executed", "not_executed"],
          f"{[i['status'] for i in job['items'][2:]]}")

    stub_reset(script="http401")
    job401 = wait_job(create(questions=[Q_TRIAD, Q_GROUND], mode="generate")["job_id"])
    check("批量·停手：认证类错误（401）同样立即停手",
          job401["state"] == "STOPPED_ON_ERROR" and job401["started_calls"] == 1
          and stub_state()["count"] == 1,
          f"state={job401['state']} started={job401['started_calls']}")

    stub_reset(script="garbage,ok")
    jobg = wait_job(create(questions=[Q_TRIAD, Q_GROUND], mode="generate")["job_id"])
    check("批量·区分错误：返回体结构异常只记这一条失败，后续继续执行",
          jobg["items"][0]["status"] == "failed" and jobg["items"][1]["status"] == "done"
          and jobg["state"] == "COMPLETED",
          f"items={[i['status'] for i in jobg['items']]} state={jobg['state']}")
    check("批量·区分错误：这类错误不触发停手（桩收到 2 次请求）",
          stub_state()["count"] == 2, f"桩计数={stub_state()['count']}")

    stub_reset(delay=0.6)
    job_c = create(questions=[Q_TRIAD, Q_GROUND, Q_FLOW], mode="generate")
    jid = job_c["job_id"]
    # 卡点必须用"桩**真的收到**了这次请求"来判定（49 号 P1 修正）。
    # 原来的写法是等 item.status == "running"，但那个状态在**检索阶段**也成立：
    # 在这个窗口里取消，按 49 号 P1 的正确行为应当是 0 请求（由阶段 4b 专门覆盖），
    # 于是"已发出的请求跑完并计数"这条判据在旧卡点下会时真时假 —— 判据不精确，
    # 不是产品行为变了。桩的 count 在收到请求、注入延迟**之前**自增，所以它是"已发出"的证据。
    check("批量·取消：取消前第 1 条的生成请求确实已发出（桩已收到 1 次请求）",
          wait_until(lambda: stub_state()["count"] == 1, timeout=10),
          f"桩计数={stub_state()['count']}")
    bj.request_cancel(jid)
    job_c = wait_job(jid)
    check("批量·取消：终态为 CANCELLED 且记录了取消位置",
          job_c["state"] == "CANCELLED" and job_c["cancelled_at_seq"] == 2,
          f"state={job_c['state']} cancelled_at={job_c['cancelled_at_seq']}")
    check("批量·取消：已发出的那一条跑完并计入实际调用数",
          job_c["items"][0]["status"] in ("done", "refused") and job_c["started_calls"] == 1,
          f"item0={job_c['items'][0]['status']} started={job_c['started_calls']}")
    check("批量·取消：之后的问题不再启动（桩只收到 1 次请求）",
          [i["status"] for i in job_c["items"][1:]] == ["not_executed", "not_executed"]
          and stub_state()["count"] == 1,
          f"桩计数={stub_state()['count']}")
    check("批量·取消：对已结束任务再次取消是幂等的（不改变结果）",
          bj.request_cancel(jid)["state"] == "CANCELLED")
    EVIDENCE["phases"].append({"name": "errors_cancel", "stub_requests": stub_state()["count"]})


# ---------------- 阶段 4b：检索期间取消（49 号 P1 定向回归） ----------------
def phase_cancel_during_retrieve() -> None:
    """在**当前题仍在检索、模型请求尚未发出**的窗口里取消：必须 0 模型请求。

    为什么单独做这一条：原有取消用例等的是"第一条生成请求已经发出"，证明的是
    "已发出的等它跑完"，并没有覆盖"检索中取消"这个窗口（49 号独立反例正是在这里
    复现出 CANCELLED 却 started_calls=1）。

    确定性做法（不靠 sleep 碰运气）：把 `rag_core.expand_and_assemble` 换成
    "进入检索 → 等放行信号"的包装函数，用事件把执行精确卡在"检索中、生成前"，
    此时调用 request_cancel()，再放行。全程只打本机桩，零真实模型请求。
    """
    import threading

    stub_reset(script="ok", delay=0.01)
    entered = threading.Event()
    release = threading.Event()
    real_assemble = rc.expand_and_assemble

    def gated_assemble(*a, **kw):
        entered.set()
        release.wait(20)
        return real_assemble(*a, **kw)

    rc.expand_and_assemble = gated_assemble      # batch_jobs 里的 rc 就是同一个模块对象
    try:
        job = create(questions=[Q_TRIAD, Q_GROUND, Q_FLOW], mode="generate")
        jid = job["job_id"]
        check("批量·取消（检索中）：第 1 题确实已进入检索阶段（窗口成立）",
              entered.wait(20), "前置条件不成立时，下面的判据没有意义")
        bj.request_cancel(jid)
        release.set()
        final = wait_job(jid)
    finally:
        rc.expand_and_assemble = real_assemble
        release.set()

    st = stub_state()
    check("批量·取消（检索中）：终态 CANCELLED 且取消位置=第 1 条",
          final["state"] == "CANCELLED" and final["cancelled_at_seq"] == 1,
          f"state={final['state']} at={final['cancelled_at_seq']}")
    check("批量·取消（检索中）：实际启动数为 0（取消之后没有补发这次生成请求）",
          final["started_calls"] == 0, f"started={final['started_calls']}")
    check("批量·取消（检索中）：桩收到 0 次请求（零模型请求是实测，不是声明）",
          st["count"] == 0, f"桩计数={st['count']}")
    check("批量·取消（检索中）：三条都如实记为未执行（不伪装成失败或空答案）",
          [i["status"] for i in final["items"]] == ["not_executed"] * 3,
          f"{[i['status'] for i in final['items']]}")
    EVIDENCE["phases"].append({"name": "cancel_during_retrieve", "stub_requests": st["count"],
                               "started_calls": final["started_calls"],
                               "state": final["state"],
                               "cancelled_at_seq": final["cancelled_at_seq"],
                               "items": [i["status"] for i in final["items"]]})


# ---------------- 阶段 5：冲突 / 重启语义 / 隔离 ----------------
def phase_conflict_and_restart() -> None:
    stub_reset(delay=0.4)
    j1 = create(questions=[Q_TRIAD, Q_GROUND, Q_FLOW], mode="generate")
    try:
        create(questions=[Q_TRIAD], mode="generate")
        check("批量·冲突：第二个生成任务被拒绝", False, "没有被拒绝")
    except bj.BatchError as exc:
        check("批量·冲突：同一进程只允许一个生成任务（generate_busy）",
              exc.code == "generate_busy", f"实际 {exc.code}")
    j2 = create(questions=[Q_TRIAD], mode="retrieve_only")
    check("批量·冲突：仅检索任务不占用生成名额（可以同时创建）",
          j2["job_id"] != j1["job_id"])
    bj.request_cancel(j1["job_id"])
    finished = wait_job(j1["job_id"])
    stub_reset(delay=0.02)

    jid = finished["job_id"]
    snapshot_file = Path(os.environ["RAG_BATCH_DIR"]) / f"{jid}.json"
    check("批量·落盘：状态快照写在隔离目录里（原子文件存在）", snapshot_file.exists(),
          snapshot_file.name)
    bj.reset()
    try:
        bj.get_job(jid)
        check("批量·重启：内存清空后旧任务不可读（404 unknown_job）", False, "仍可读")
    except bj.BatchError as exc:
        check("批量·重启：内存清空后旧任务不可读，也绝不自动重放",
              exc.code == "unknown_job", f"实际 {exc.code}")
    check("批量·重启：磁盘上残留的快照只是证据文件，不会被自动执行",
          snapshot_file.exists() and bj.active_generate_jobs() == [])
    check("批量·隔离：状态目录在运行目录内，不是内置资料目录或用户资料目录",
          "batches" in str(bj.BATCH_DIR)
          and str(REPO / "data") not in str(bj.BATCH_DIR)
          and str(REPO / "用户资料") not in str(bj.BATCH_DIR))
    EVIDENCE["phases"].append({"name": "conflict_restart", "batch_dir": str(bj.BATCH_DIR)})


# ---------------- 阶段 6：导出 ----------------
NASTY = [
    "逗号,分隔的问题？",
    "双引号\"引号\"的问题？",
    "带换行\n第二行的问题？",
    "../../etc/passwd 路径穿越的问题？",
    "<script>alert(1)</script> 类 HTML 的问题？",
    "**markdown** [链接](http://example.invalid) 的问题？",
    "=cmd|' /C calc'!A1 公式注入形态的问题？",
    "制表\t符与中文，全角逗号的问题？",
    "- 以减号开头的行？",
    "@ 以 at 开头的问题？",
]


def _csv_expect(job: dict) -> list:
    """从任务快照独立复算 CSV 应当长什么样（15 列：前 10 列引用语义 + 后 5 列命中证据）。"""
    rows = []
    for it in job["items"]:
        refs = it.get("refs") or []
        hits = it.get("hits") or []
        rows.append([
            str(it["sequence"]), it["question"], it["status"], it.get("answer") or "",
            json.dumps([r["id"] for r in refs], ensure_ascii=False),
            json.dumps([r["source_title"] for r in refs], ensure_ascii=False),
            it.get("error") or "", job["mode"], job["corpus_name"], str(job["top_k"]),
            json.dumps([h.get("id") or "" for h in hits], ensure_ascii=False),
            json.dumps([h.get("source_title") or "" for h in hits], ensure_ascii=False),
            json.dumps([h.get("section") or "" for h in hits], ensure_ascii=False),
            json.dumps([h.get("score") for h in hits], ensure_ascii=False),
            json.dumps([h.get("text") or "" for h in hits], ensure_ascii=False),
        ])
    return rows


def _hits_from_row(row: list) -> list:
    """把 CSV 一行里的五个 hit_* 列**还原成结构化命中**（顺序即数组顺序）。"""
    ids, titles, sections, scores, snippets = (json.loads(row[i]) for i in range(10, 15))
    return [{"id": ids[k], "source_title": titles[k], "section": sections[k],
             "score": scores[k], "text": snippets[k]} for k in range(len(ids))]


def _hits_want(item: dict) -> list:
    """同一结构的"应有值"，来自任务快照（与导出侧各写各的，避免自证）。"""
    return [{"id": h.get("id") or "", "source_title": h.get("source_title") or "",
             "section": h.get("section") or "", "score": h.get("score"),
             "text": h.get("text") or ""} for h in (item.get("hits") or [])]


def _docx_text(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)


def _docx_part(data: bytes, name: str) -> str:
    """读 DOCX 包内某个 part 的原文（结构级检查用：字体声明、东亚语言、XML 完整性）。"""
    import zipfile
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(name).decode("utf-8")


def _docx_parts(data: bytes) -> dict:
    """把包里所有 xml/rels part 解出来（顺带证明 zip 与每个 XML 都完整可解析）。"""
    import zipfile
    from xml.etree import ElementTree
    out: dict = {}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        bad = z.testzip()
        if bad is not None:
            raise AssertionError(f"zip 内容损坏：{bad}")
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                raw = z.read(name)
                ElementTree.fromstring(raw)          # 解析失败直接抛错
                out[name] = raw.decode("utf-8")
    return out


def check_csv_out(job: dict, prefix: str) -> bytes:
    data = bx.render_csv(job)
    check(f"{prefix}·CSV：以 UTF-8 BOM 开头（Excel 打开中文不乱码）",
          data.startswith(b"\xef\xbb\xbf"))
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    check(f"{prefix}·CSV：表头就是约定的 15 列（10 列引用语义 + 5 列命中证据）",
          tuple(rows[0]) == bx.CSV_HEADERS, f"实际 {rows[0]}")
    check(f"{prefix}·CSV：每行的五个 hit_* 列还原后与任务快照的 hits 逐字段一致",
          all(_hits_from_row(row) == _hits_want(it)
              for row, it in zip(rows[1:], job["items"])),
          f"首行还原={_hits_from_row(rows[1])[:1] if len(rows) > 1 else None}")
    check(f"{prefix}·CSV：行数 = 题目数 + 1（顺序即输入顺序）",
          len(rows) == job["total"] + 1, f"rows={len(rows)} total={job['total']}")
    check(f"{prefix}·CSV：逐字段与任务记录一致（问题/状态/答案/引用/错误/参数）",
          rows[1:] == _csv_expect(job))
    return data


def check_docx_out(job: dict, prefix: str) -> bytes:
    data = bx.render_docx(job)
    check(f"{prefix}·DOCX：是合法的 docx（PK 头 + python-docx 能打开）", data[:2] == b"PK")
    text = _docx_text(data)
    missing = [it["sequence"] for it in job["items"] if it["question"] not in text]
    check(f"{prefix}·DOCX：每道问题原文都在报告里", not missing, f"缺失 {missing}")
    check(f"{prefix}·DOCX：含固定边界说明（引用有效不等于回答正确）",
          bx.DOCX_BOUNDARY in text)
    check(f"{prefix}·DOCX：含知识库名称、运行模式、Top-K、问题总数",
          job["corpus_name"] in text and str(job["top_k"]) in text
          and f"问题总数：{job['total']}" in text)
    check(f"{prefix}·DOCX：含计划/实际模型请求数与零重试事实",
          f"计划模型请求数：{job['planned_calls']}" in text
          and f"实际启动：{job['started_calls']}" in text and "自动重试：0" in text)
    refs_ok = True
    for it in job["items"]:
        for r in it["refs"]:
            if f"[{r['id']}]" not in text or r["source_title"] not in text:
                refs_ok = False
    check(f"{prefix}·DOCX：引用 id 与来源标题逐条出现", refs_ok)
    return data


def phase_export() -> None:
    # 四类任务：完成 / 拒答 / 失败(停手) / 取消后部分完成
    stub_reset(delay=0.02)
    job_done = wait_job(create(questions=[Q_TRIAD, Q_GROUND, Q_FLOW], mode="generate")["job_id"])

    stub_reset(script="refuse")
    job_refuse = wait_job(create(questions=[Q_TRIAD, Q_GROUND], mode="generate")["job_id"])
    check("导出·拒答任务：两条都记为 refused（不混进完成计数）",
          [i["status"] for i in job_refuse["items"]] == ["refused", "refused"]
          and job_refuse["refused"] == 2, f"{[i['status'] for i in job_refuse['items']]}")

    stub_reset(script="ok,http500")
    job_fail = wait_job(create(questions=[Q_TRIAD, Q_GROUND], mode="generate")["job_id"])
    check("导出·失败任务：一条 done、一条 failed、任务为 STOPPED_ON_ERROR",
          job_fail["items"][0]["status"] == "done"
          and job_fail["items"][1]["status"] == "failed"
          and job_fail["state"] == "STOPPED_ON_ERROR",
          f"{[i['status'] for i in job_fail['items']]} / {job_fail['state']}")

    stub_reset(delay=0.6)
    job_cancel = create(questions=[Q_TRIAD, Q_GROUND, Q_FLOW], mode="generate")
    # 同阶段 4 的口径：用"桩已收到请求"卡点，样本才是稳定的"取消后部分完成"
    # （49 号 P1 之前这里等的是 item.status=running，可能落在检索窗口里 → 样本会变成全未执行）
    wait_until(lambda: stub_state()["count"] == 1, timeout=10)
    bj.request_cancel(job_cancel["job_id"])
    job_cancel = wait_job(job_cancel["job_id"])
    stub_reset(delay=0.02)

    stub_before = stub_state()["count"]
    started_before = job_done["started_calls"]
    check("导出·四类任务都能导出：完成 / 拒答 / 失败(停手) / 取消后部分完成",
          all(j["state"] in bj.FINAL_STATES for j in (job_done, job_refuse, job_fail, job_cancel)))

    for job, prefix in ((job_done, "完成"), (job_refuse, "拒答"),
                        (job_fail, "失败"), (job_cancel, "取消")):
        check_csv_out(job, prefix)
        check_docx_out(job, prefix)

    check("导出·安全：导出动作零新增调用（桩计数与任务台账都不变）",
          stub_state()["count"] == stub_before and job_done["started_calls"] == started_before,
          f"桩 {stub_before}→{stub_state()['count']}")

    job_nasty = wait_job(create(questions=NASTY, mode="retrieve_only")["job_id"])
    data_nasty = bx.render_csv(job_nasty)
    rows = list(csv.reader(io.StringIO(data_nasty.decode("utf-8-sig"))))
    check("导出·反例：逗号/换行/引号/制表/Markdown/类 HTML 逐字节还原（CSV 转义正确）",
          [r[1] for r in rows[1:]] == NASTY,
          f"实际前三条={[r[1] for r in rows[1:]][:3]}")
    text_nasty = _docx_text(bx.render_docx(job_nasty))
    check("导出·反例：同样的问题在 DOCX 里也完整保留",
          all(q in text_nasty for q in NASTY))
    name_csv = bx.export_filename(job_nasty, "csv")
    name_docx = bx.export_filename(job_nasty, "docx")
    check("导出·反例：文件名只由固定前缀 + 时间戳构成（用户文本不参与路径）",
          re.fullmatch(r"RAG批量问答报告-\d{8}-\d{6}\.csv", name_csv) is not None
          and re.fullmatch(r"RAG批量问答报告-\d{8}-\d{6}\.docx", name_docx) is not None,
          f"{name_csv} / {name_docx}")

    text_by_case = {
        "完成": (bx.render_csv(job_done).decode("utf-8-sig"), _docx_text(bx.render_docx(job_done))),
        "失败": (bx.render_csv(job_fail).decode("utf-8-sig"), _docx_text(bx.render_docx(job_fail))),
        "取消": (bx.render_csv(job_cancel).decode("utf-8-sig"),
                 _docx_text(bx.render_docx(job_cancel))),
    }
    leak = [f"{case}/{kind}" for case, pair in text_by_case.items()
            for kind, blob in (("csv", pair[0]), ("docx", pair[1]))
            if SENTINEL_KEY in blob or "Authorization" in blob or "Bearer " in blob]
    check("导出·安全：报告里都不含哨兵 Key、不含 Authorization/Bearer", not leak,
          f"命中 {leak}")
    # 负向环视排除 URL 里的 "s://"，与 verify_flows.py 的 S10 扫描同一口径
    abspath = [f"{case}/{kind}" for case, pair in text_by_case.items()
               for kind, blob in (("csv", pair[0]), ("docx", pair[1]))
               if re.search(r"(?<![A-Za-z])[A-Za-z]:[\\/]", blob)]
    check("导出·安全：报告里不含开发机绝对路径", not abspath, f"命中 {abspath}")

    _EXPORT_JOB_IDS.update({"done": job_done["job_id"], "cancel": job_cancel["job_id"]})
    EVIDENCE["export_samples"] = {
        "csv_done": bx.render_csv(job_done).decode("utf-8-sig"),
        "csv_cancel": bx.render_csv(job_cancel).decode("utf-8-sig"),
        "docx_bytes": len(bx.render_docx(job_done)),
        "summary_done": bj.summary(job_done),
        "summary_fail": bj.summary(job_fail),
        "summary_cancel": bj.summary(job_cancel),
    }
    EVIDENCE["phases"].append({"name": "export", "stub_requests": stub_state()["count"]})


# ---------------- 阶段 6b：命中证据落地 / 中文字体（50 号工单五节） ----------------
def _hits_docx_missing(dtext: str, job: dict) -> list:
    """DOCX 里逐条核命中区块：排名 / id / 来源 / 章节 / 分数 / 原文（按截断长度）是否都在。"""
    miss = []
    for it in job["items"]:
        hits = it.get("hits") or []
        for rank, h in enumerate(hits, 1):
            want = [f"命中 {rank}／{len(hits)}：{h.get('id') or ''}",
                    f"来源标题：{h.get('source_title') or '（无）'}",
                    f"章节：{h.get('section') or '（无章节标题）'}",
                    f"BM25 分数：{h.get('score')}"]
            text = (h.get("text") or "")[:bx.SNIPPET_MAX_CHARS]
            if any(w not in dtext for w in want) or (text and text not in dtext):
                miss.append(f"第{it['sequence']}题#{rank}")
    return miss


def phase_hit_evidence() -> None:
    """工单五节的定向回归：命中可自证、两类引用不互相冒充、导出零新增调用、中文字体落地。

    全程只打本机桩（仅检索任务本身不发任何请求，生成样本也都指向本机桩）。
    """
    # ---- 1) 仅检索任务：命中快照与报告逐字段一致（工单 1 / 2）----
    stub_reset()
    job = wait_job(create(questions=[Q_TRIAD, Q_GROUND, Q_FLOW],
                          mode="retrieve_only")["job_id"])
    item0 = job["items"][0]
    check("工单1·仅检索：命中随任务一起保存了原文（不再只有 chunk id）",
          bool(item0["hits"]) and all(h.get("text") for h in item0["hits"]),
          f"hits={len(item0['hits'])} 带原文={sum(1 for h in item0['hits'] if h.get('text'))}")
    exp = rc.expand_and_assemble(Q_TRIAD, top_k=3, corpus_id=rc.RAG_LEARNING)
    check("工单1·仅检索：id / 顺序 / 来源 / 章节 / 分数与直接调 rag_core 逐字段一致",
          [(h["id"], h["source_title"], h["section"], h["score"]) for h in item0["hits"]]
          == [(e["id"], e["source_title"], e["section"], e["score"]) for e in exp["hits"]])
    docs = {d["id"]: d for d in rc.load_documents(rc.RAG_LEARNING)}
    drifted = [h["id"] for it in job["items"] for h in it["hits"]
               if docs.get(h["id"], {}).get("text") != h.get("text")]
    check("工单1·仅检索：保存的命中原文与语料库当前 chunk 文本逐字一致（不靠导出侧自证）",
          not drifted, f"对不上 {drifted}")

    csv_bytes = bx.render_csv(job)
    rows = list(csv.reader(io.StringIO(csv_bytes.decode("utf-8-sig"))))
    check("工单1·仅检索：CSV 命中列还原后与快照完全一致（含原文，顺序一致）",
          [_hits_from_row(r) for r in rows[1:]] == [_hits_want(i) for i in job["items"]])
    # 注意：这里必须查到**元素**一级 —— 只看"数组非空"会漏过 [""] 这种
    # "列里有值但值是空串"的情况（负向探针 C 抓到过这个弱点）。
    def _hit_columns_filled(rows_: list) -> bool:
        for row in rows_:
            for idx in range(10, 15):
                vals = json.loads(row[idx])
                if not vals or any(v is None or not str(v).strip() for v in vals):
                    return False
        return True

    check("工单2·仅检索：citation_ids 允许为空，但五个 hit_* 列每行每个元素都非空",
          all(json.loads(r[4]) == [] for r in rows[1:]) and _hit_columns_filled(rows[1:]),
          f"每行命中数={[len(json.loads(r[10])) for r in rows[1:]]}")
    check("工单2·仅检索：命中原文列不是空串占位（每个 snippet 与快照逐字相同）",
          all(json.loads(r[14]) == [h.get("text") or "" for h in it["hits"]]
              and all(s.strip() for s in json.loads(r[14]))
              for r, it in zip(rows[1:], job["items"])))

    # ---- 2) DOCX 逐题结构与措辞（工单 3）----
    docx_bytes = bx.render_docx(job)
    dtext = _docx_text(docx_bytes)
    check("工单3·仅检索：不再出现「（无回答文本）」「引用：无」这类易被读成「什么都没找到」的措辞",
          "（无回答文本）" not in dtext and "引用：无" not in dtext)
    check("工单3·仅检索：改为「回答：未生成（仅检索模式）」+「回答引用：不适用（未生成回答）」",
          "回答：未生成（仅检索模式）" in dtext
          and "回答引用：不适用（未生成回答）" in dtext)
    miss = _hits_docx_missing(dtext, job)
    check("工单3·仅检索：DOCX 每条命中独立成块（排名/id/来源/章节/BM25 分数/原文都在）",
          not miss, f"缺失 {miss[:6]}")
    long_hits = [h for it in job["items"] for h in it["hits"]
                 if len(h.get("text") or "") > bx.SNIPPET_MAX_CHARS]
    check("工单3·仅检索：超长命中明确标注截断并指向 CSV（不让读者误以为原文就这么短）",
          (all("此处截断显示前" in dtext and "完整文本见 CSV 的 hit_snippets 列" in dtext
               for _ in long_hits) if long_hits else True),
          f"超长命中 {len(long_hits)} 条")
    from docx import Document
    parsed = Document(io.BytesIO(docx_bytes))
    heads = [p for p in parsed.paragraphs if p.style.name == "Heading 2"]
    qparas = [p for p in parsed.paragraphs if p.text.startswith("问题：")]
    check("工单3·仅检索：题号标题与问题都设了「与下段同页」（不会把题号孤悬在页尾）",
          bool(heads) and bool(qparas)
          and all(p.paragraph_format.keep_with_next for p in heads)
          and all(p.paragraph_format.keep_with_next for p in qparas),
          f"标题 {len(heads)} 个 / 问题 {len(qparas)} 条")

    # ---- 3) 生成任务：回答引用与检索命中分开验证（工单 3 后半）----
    stub_reset(delay=0.02)
    job_g = wait_job(create(questions=[Q_TRIAD, Q_GROUND], mode="generate")["job_id"])
    rows_g = list(csv.reader(io.StringIO(bx.render_csv(job_g).decode("utf-8-sig"))))
    check("工单3·生成：citation 列只装回答引用、hit 列只装检索命中（各取各的快照）",
          all(json.loads(r[4]) == [x["id"] for x in it["refs"]]
              and json.loads(r[10]) == [h["id"] for h in it["hits"]]
              for r, it in zip(rows_g[1:], job_g["items"])))
    stub_reset(script="refuse")
    job_r = wait_job(create(questions=[Q_TRIAD, Q_GROUND], mode="generate")["job_id"])
    rows_r = list(csv.reader(io.StringIO(bx.render_csv(job_r).decode("utf-8-sig"))))
    check("工单3·生成：模型没引用任何片段时 citation 列为空、hit 列照常非空（不互相冒充）",
          all(it["refs"] == [] for it in job_r["items"])
          and all(json.loads(r[4]) == [] for r in rows_r[1:])
          and all(json.loads(r[10]) for r in rows_r[1:]),
          f"citations={[json.loads(r[4]) for r in rows_r[1:]]} "
          f"命中数={[len(json.loads(r[10])) for r in rows_r[1:]]}")

    # ---- 4) 导出零新增调用、不改快照（工单 4）----
    stub_reset()
    job_e = wait_job(create(questions=[Q_TRIAD, Q_GROUND], mode="retrieve_only")["job_id"])
    snapshot_before = json.dumps(job_e, ensure_ascii=False, sort_keys=True)
    csv1, docx1 = bx.render_csv(job_e), bx.render_docx(job_e)
    csv2, docx2 = bx.render_csv(job_e), bx.render_docx(job_e)
    check("工单4·导出：反复导出（CSV/DOCX 各两次）不增加任何模型请求",
          stub_state()["count"] == 0 and job_e["started_calls"] == 0,
          f"桩计数={stub_state()['count']}")
    check("工单4·导出：导出不改动任务快照（导出前后逐字段一致）",
          json.dumps(job_e, ensure_ascii=False, sort_keys=True) == snapshot_before)
    check("工单4·导出：同一快照两次渲染 CSV 逐字节一致（导出的确是纯读取）", csv1 == csv2)
    check("工单4·导出：DOCX 两次渲染都能被 python-docx 打开（生成失败也不会留半截）",
          docx1[:2] == b"PK" and docx2[:2] == b"PK")

    # ---- 5) 中文字体：结构检查（工单 5）+ 中文样本（工单 6）----
    parts = _docx_parts(docx1)
    check("工单7·完整性：DOCX 包内每个 XML/rels part 都能解析（zip + XML 完整性）",
          {"word/document.xml", "word/styles.xml", "word/settings.xml"} <= set(parts),
          f"parts={len(parts)}")
    doc_xml, styles_xml = parts["word/document.xml"], parts["word/styles.xml"]
    # 判据口径按工单原文："只要**已使用样式和正文 runs** 有明确东亚字体即可"。
    # 未使用的内置样式（Heading 3–9、Quote、List Paragraph…）不参与渲染，不做要求，
    # 否则脚本会为了"清零"去重写整张样式表 —— 那超出工单范围，也不是缺陷。
    THEME_RE = r'w:(?:ascii|hAnsi|eastAsia|cs)Theme="[^"]+"'
    doc_themed = re.findall(THEME_RE, doc_xml)
    check("工单5·字体：正文（document.xml）里 0 处主题字体引用（不再由主题猜字体）",
          not doc_themed, f"残留 {len(doc_themed)} 处")
    dflt = re.search(r"<w:docDefaults>.*?</w:docDefaults>", styles_xml, re.S)
    dflt_block = dflt.group(0) if dflt else ""
    check("工单5·字体：docDefaults 写死东亚字体、且不含主题引用",
          f'w:eastAsia="{bx.CJK_FONT}"' in dflt_block
          and not re.findall(THEME_RE, dflt_block),
          f"docDefaults 命中={bool(dflt)}")
    style_map = {}
    for block in re.findall(r"<w:style\b.*?</w:style>", styles_xml, re.S):
        sid = re.search(r'w:styleId="([^"]+)"', block)
        if sid:
            style_map[sid.group(1)] = block
    # 用 python-docx 解析出的 styleId 定位样式元素，而不是拿界面名去猜 w:name
    # （Word 内置样式的 w:name 是小写的 "heading 1"，按名字匹配会假 FAIL）。
    used_styles = ("Normal", "Title", "Heading 1", "Heading 2")
    style_ids = {}
    for name in used_styles:
        try:
            style_ids[name] = parsed.styles[name].style_id
        except KeyError:
            style_ids[name] = None
    missing = [n for n, sid in style_ids.items()
               if not sid or f'w:eastAsia="{bx.CJK_FONT}"' not in style_map.get(sid, "")]
    still_themed = [n for n, sid in style_ids.items()
                    if sid and re.findall(THEME_RE, style_map.get(sid, ""))]
    check("工单5·字体：已使用的 4 个样式都写死东亚字体且清掉主题引用",
          not missing and not still_themed,
          f"缺声明 {missing}；仍带主题引用 {still_themed}；styleId={style_ids}")
    runs = re.findall(r"<w:r(?:\s[^>]*)?>.*?</w:r>", doc_xml, re.S)
    bare = [i for i, r in enumerate(runs) if f'w:eastAsia="{bx.CJK_FONT}"' not in r]
    check("工单5·字体：document.xml 里每个 run 都有显式 w:eastAsia（兜底层真的跑了）",
          bool(runs) and not bare, f"runs={len(runs)} 未声明={len(bare)}")
    lang = re.search(r"<w:themeFontLang[^>]*/?>", parts["word/settings.xml"])
    check("工单5·字体：settings.xml 的 themeFontLang@eastAsia 已设为 zh-CN（不再是 ja-JP）",
          lang is not None and 'w:eastAsia="zh-CN"' in lang.group(0),
          lang.group(0) if lang else "缺少该元素")

    sample_chars = "项目试导测资料标题为习仅务动"
    qs_cjk = [f"{i}．{sample_chars} 第 {i} 题的中文样本" for i in range(1, 21)]
    job_cjk = wait_job(create(questions=qs_cjk, mode="retrieve_only")["job_id"])
    docx_cjk = bx.render_docx(job_cjk)
    dtext_cjk = _docx_text(docx_cjk)
    check("工单6·中文样本：20 题里每个中文字符都原样进了报告（无丢字 / 无占位）",
          all(c in dtext_cjk for c in sample_chars)
          and all(q in dtext_cjk for q in qs_cjk))
    check("工单6·中文样本：成包里没有方框字符 U+25A1、也没有替换符 U+FFFD",
          "\u25a1" not in dtext_cjk and "\ufffd" not in dtext_cjk)
    _EXTRA_FILES["sample_cjk_20q.docx"] = docx_cjk
    _EXTRA_FILES["sample_cjk_20q.csv"] = bx.render_csv(job_cjk)

    # ---- 6) 安全：凭据 / 授权头 / 开发机绝对路径（工单 7）----
    csv1_text = csv1.decode("utf-8-sig")
    leak = [name for name, text in parts.items()
            if SENTINEL_KEY in text or "Authorization" in text or "Bearer " in text]
    if SENTINEL_KEY in csv1_text or "Authorization" in csv1_text or "Bearer " in csv1_text:
        leak.append("csv")
    check("工单7·安全：命中证据列没有把凭据 / Authorization 写进报告（含包内 XML）",
          not leak, f"命中 {leak}")
    abspath = [name for name, text in parts.items()
               if re.search(r"(?<![A-Za-z])[A-Za-z]:[\\/]", text)]
    if re.search(r"(?<![A-Za-z])[A-Za-z]:[\\/]", csv1_text):
        abspath.append("csv")
    check("工单7·安全：报告里没有开发机绝对路径（含包内 XML）", not abspath, f"命中 {abspath}")
    check("工单7·安全：CSV 仍是 UTF-8 BOM（Excel 打开中文不乱码）",
          csv1.startswith(b"\xef\xbb\xbf"))
    _EXTRA_FILES["sample_hit_evidence.docx"] = docx1
    _EXTRA_FILES["sample_hit_evidence.csv"] = csv1
    EVIDENCE["phases"].append({"name": "hit_evidence", "stub_requests": stub_state()["count"],
                               "retrieve_items": job["total"],
                               "runs_with_eastasia": len(runs) - len(bare),
                               "docx_bytes": len(docx1)})


# ---------------- 阶段 7：HTTP 契约（真跑 api_v3） ----------------
def phase_http() -> None:
    # --- 无 Key 的 api_v3：契约与拒绝路径 ---
    api = start_api(with_key=False)
    try:
        code, _, body, _ = http("GET", "/api/meta")
        meta = jbody(body)
        batch = meta.get("batch") or {}
        check("接口·批量：meta 给出服务端上限（1–20、串行、零重试）",
              code == 200 and batch.get("min_questions") == 1 and batch.get("max_questions") == 20
              and batch.get("concurrency") == 1 and batch.get("auto_retry") == 0,
              f"batch={batch}")
        check("接口·批量：meta 的导出格式就是 csv / docx（不做 PDF）",
              batch.get("export_formats") == ["csv", "docx"], f"{batch.get('export_formats')}")

        payload = {"corpus_id": rc.RAG_LEARNING, "mode": "retrieve_only", "top_k": 3,
                   "questions": [Q_TRIAD, Q_GROUND], "confirm_calls": 0}
        code, _, body, _ = http("POST", "/api/batch-jobs", payload)
        job = jbody(body).get("job") or {}
        jid = job.get("job_id", "")
        check("接口·批量：创建检索任务返回 200 + job_id + summary",
              code == 200 and jid and jbody(body).get("summary", {}).get("total") == 2,
              f"code={code} id={jid[:8]}")

        code, _, body, _ = http("GET", f"/api/batch-jobs/{jid}")
        check("接口·批量：按 job_id 能取回进度与结果",
              code == 200 and jbody(body).get("job", {}).get("job_id") == jid)

        code, _, body, _ = http("GET", "/api/batch-jobs/deadbeefdeadbeefdeadbeef")
        check("接口·批量：未知 job_id → 404/unknown_job",
              code == 404 and jbody(body).get("error", {}).get("code") == "unknown_job",
              f"{code}/{jbody(body).get('error', {}).get('code')}")

        code, _, body, _ = http("POST", "/api/batch-jobs",
                                {"corpus_id": rc.RAG_LEARNING, "mode": "retrieve_only", "top_k": 3})
        check("接口·批量：缺 questions → 400/missing_field",
              code == 400 and jbody(body).get("error", {}).get("code") == "missing_field",
              f"{code}/{jbody(body).get('error', {}).get('code')}")

        code, _, body, _ = http("POST", "/api/batch-jobs", {**payload, "extra": 1})
        check("接口·批量：多带字段 → 400/unexpected_field",
              code == 400 and jbody(body).get("error", {}).get("code") == "unexpected_field",
              f"{code}/{jbody(body).get('error', {}).get('code')}")

        code, _, body, _ = http("POST", "/api/batch-jobs",
                                {"corpus_id": rc.RAG_LEARNING, "mode": "generate", "top_k": 3,
                                 "questions": [Q_TRIAD], "confirm_calls": 1})
        check("接口·批量：无 Key 时生成模式在建任务阶段就被拒（409/no_api_key，零请求）",
              code == 409 and jbody(body).get("error", {}).get("code") == "no_api_key",
              f"{code}/{jbody(body).get('error', {}).get('code')}")

        code, _, body, _ = http("POST", "/api/batch-jobs",
                                {"corpus_id": "user_library", "mode": "retrieve_only", "top_k": 3,
                                 "questions": [Q_TRIAD], "confirm_calls": 0})
        check("接口·批量：空库 fail closed（409/user_library_empty）",
              code == 409 and jbody(body).get("error", {}).get("code") == "user_library_empty",
              f"{code}/{jbody(body).get('error', {}).get('code')}")

        code, _, _, _ = http("GET", "/api/batch-jobs")
        check("接口·批量：GET /api/batch-jobs → 405（只接受 POST）", code == 405, f"code={code}")

        code, _, body, _ = http("POST", f"/api/batch-jobs/{jid}/cancel")
        check("接口·批量：取消端点返回 200 且带回任务状态",
              code == 200 and jbody(body).get("job", {}).get("state") in
              ("PENDING", "RUNNING", "COMPLETED", "CANCELLED"), f"code={code}")

        code, _, body, _ = http("GET", f"/api/batch-jobs/{jid}/export?format=pdf")
        check("接口·批量：不支持的导出格式 → 400/unknown_format",
              code == 400 and jbody(body).get("error", {}).get("code") == "unknown_format",
              f"{code}/{jbody(body).get('error', {}).get('code')}")

        code, headers, _, raw = http("GET", f"/api/batch-jobs/{jid}/export?format=csv")
        check("接口·批量：CSV 导出是 200 + text/csv + 附件文件名（含 RFC 5987 中文名）",
              code == 200 and "text/csv" in (headers.get("Content-Type") or "")
              and "attachment" in (headers.get("Content-Disposition") or "")
              and "filename*=UTF-8''" in (headers.get("Content-Disposition") or ""),
              f"code={code} disp={(headers.get('Content-Disposition') or '')[:60]}")
        _, _, job_body, _ = http("GET", f"/api/batch-jobs/{jid}")
        job_http = jbody(job_body).get("job") or {}
        check("接口·批量：CSV 导出内容与该 job 的接口记录逐字节一致",
              raw == bx.render_csv(job_http), f"bytes={len(raw)}")

        code, headers, _, raw = http("GET", f"/api/batch-jobs/{jid}/export?format=docx")
        check("接口·批量：DOCX 导出是 200 + docx 内容类型 + 合法 PK 文件",
              code == 200 and "wordprocessingml" in (headers.get("Content-Type") or "")
              and raw[:2] == b"PK", f"code={code}")
    finally:
        stop_process(api)

    # --- 有 Key 的 api_v3：端到端（仍然只打本机桩）---
    api = start_api(with_key=True)
    try:
        stub_reset(delay=0.02)
        code, _, body, _ = http("POST", "/api/batch-jobs",
                                {"corpus_id": rc.RAG_LEARNING, "mode": "generate", "top_k": 3,
                                 "questions": [Q_TRIAD, Q_GROUND], "confirm_calls": 2})
        job = jbody(body).get("job") or {}
        jid = job.get("job_id", "")
        check("接口·批量：有 Key 时创建生成任务成功（计划数 2）",
              code == 200 and jid and job.get("planned_calls") == 2, f"code={code}")
        final: dict = {}
        for _ in range(600):
            _, _, body, _ = http("GET", f"/api/batch-jobs/{jid}")
            final = jbody(body).get("job") or {}
            if final.get("state") in bj.FINAL_STATES:
                break
            time.sleep(0.1)
        check("接口·批量：端到端跑完（COMPLETED、实际启动 2、零重试）",
              final.get("state") == "COMPLETED" and final.get("started_calls") == 2
              and final.get("retries") == 0,
              f"state={final.get('state')} started={final.get('started_calls')}")
        check("接口·批量：端到端只打本机桩（桩计数 2），没有真实模型请求",
              stub_state()["count"] == 2, f"桩计数={stub_state()['count']}")
        code, _, _, raw = http("GET", f"/api/batch-jobs/{jid}/export?format=docx")
        check("接口·批量：端到端任务的 DOCX 导出可用且含边界说明",
              code == 200 and bx.DOCX_BOUNDARY in _docx_text(raw), f"code={code}")
        EVIDENCE["http_generate"] = {"state": final.get("state"),
                                     "started_calls": final.get("started_calls"),
                                     "retries": final.get("retries")}
    finally:
        stop_process(api)


# ---------------- main ----------------
def main() -> int:
    import hashlib

    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="")
    ap.add_argument("--keep", action="store_true", help="保留临时隔离目录（默认删除）")
    args = ap.parse_args()
    out_dir = Path(args.out) if args.out else (WORK / "evidence")
    out_dir.mkdir(parents=True, exist_ok=True)

    global MANIFEST_HASH_BEFORE
    MANIFEST_HASH_BEFORE = hashlib.sha256(MANIFEST.read_bytes()).hexdigest()

    start_stub(delay=0.02)
    try:
        print("--- 阶段 1：输入校验（fail closed）---")
        phase_validation()
        print("\n--- 阶段 2：仅检索模式（模型请求必须为 0）---")
        phase_retrieve()
        print("\n--- 阶段 3：生成模式（1–20 条硬上限、串行、零重试）---")
        phase_generate()
        print("\n--- 阶段 4：系统性错误停手 / 非系统错误 / 取消 ---")
        phase_errors_and_cancel()
        print("\n--- 阶段 4b：检索期间取消（49 号 P1：必须 0 模型请求）---")
        phase_cancel_during_retrieve()
        print("\n--- 阶段 5：并发冲突 / 重启语义 / 隔离 ---")
        phase_conflict_and_restart()
        print("\n--- 阶段 6：DOCX / CSV 导出与反例 ---")
        phase_export()
        print("\n--- 阶段 6b：命中证据落地 / 中文字体（50 号工单五节）---")
        phase_hit_evidence()
        print("\n--- 阶段 7：薄入口 HTTP 契约（真跑 api_v3，仍只打本机桩）---")
        phase_http()
    finally:
        for proc in reversed(PROCESSES):
            stop_process(proc)

    manifest_after = hashlib.sha256(MANIFEST.read_bytes()).hexdigest()
    check("隔离：内置语料 manifest 在整轮验收中逐字节未变",
          manifest_after == MANIFEST_HASH_BEFORE,
          f"{MANIFEST_HASH_BEFORE[:16]} → {manifest_after[:16]}")
    api_log = WORK / "api_with_key.log"
    log_text = api_log.read_text(encoding="utf-8", errors="ignore") if api_log.exists() else ""
    check("全程：日志里没有真实模型域名（生成目标只有本机桩）",
          "api.deepseek.com" not in log_text and f"127.0.0.1:{STUB_PORT}" in log_text)

    result = {"service": "batch_jobs+batch_export", "passed": len(PASSED), "failed": len(FAILED),
              "passed_names": PASSED, "failed_names": FAILED, "evidence": EVIDENCE,
              "batch_dir": str(bj.BATCH_DIR), "work_dir": str(WORK)}
    (out_dir / "batch_jobs_result.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    samples = EVIDENCE.get("export_samples") or {}
    if samples.get("csv_done"):
        (out_dir / "sample_report_done.csv").write_text(samples["csv_done"], encoding="utf-8-sig")
    if samples.get("csv_cancel"):
        (out_dir / "sample_report_cancel.csv").write_text(samples["csv_cancel"],
                                                          encoding="utf-8-sig")
    for fname, blob in _EXTRA_FILES.items():          # 50 号工单：命中证据 / 中文样本
        (out_dir / fname).write_bytes(blob)
        print(f"[样本] {fname}（{len(blob)} 字节）")

    print("\n" + "=" * 70)
    if FAILED:
        print("FAIL 详情：")
        for name in FAILED:
            print(f"  - {name}")
    print(f"批量与导出验收：PASS {len(PASSED)} / FAIL {len(FAILED)}")
    print(f"证据写入：{out_dir}")
    if not args.keep:
        shutil.rmtree(WORK, ignore_errors=True)
        print(f"[verify_batch_jobs] 临时隔离目录已清理：{WORK}")
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())




